"""
Tests for MOU and Invoice generation tool.

Tests the template population logic, school year calculation,
and document generation functionality.
"""

import os
import re
from datetime import datetime, timedelta
from pathlib import Path
from unittest.mock import AsyncMock, MagicMock, patch

import pytest


# =============================================================================
# Unit Tests for Helper Functions
# =============================================================================

class TestCalculateSchoolYear:
    """Tests for _calculate_school_year function."""

    def test_school_year_in_fall_semester(self):
        """Test school year calculation when month is Sept-Dec."""
        from src.tools.finance import _calculate_school_year

        # Mock datetime to return October 2025
        with patch('src.tools.finance.datetime') as mock_datetime:
            mock_datetime.now.return_value = datetime(2025, 10, 15)
            result = _calculate_school_year()
            assert result == "2025-2026"

    def test_school_year_in_spring_semester(self):
        """Test school year calculation when month is Jan-Aug."""
        from src.tools.finance import _calculate_school_year

        # Mock datetime to return March 2026
        with patch('src.tools.finance.datetime') as mock_datetime:
            mock_datetime.now.return_value = datetime(2026, 3, 15)
            result = _calculate_school_year()
            assert result == "2025-2026"

    def test_school_year_boundary_august(self):
        """Test school year at August (still spring semester)."""
        from src.tools.finance import _calculate_school_year

        with patch('src.tools.finance.datetime') as mock_datetime:
            mock_datetime.now.return_value = datetime(2025, 8, 31)
            result = _calculate_school_year()
            assert result == "2024-2025"

    def test_school_year_boundary_september(self):
        """Test school year at September (new fall semester)."""
        from src.tools.finance import _calculate_school_year

        with patch('src.tools.finance.datetime') as mock_datetime:
            mock_datetime.now.return_value = datetime(2025, 9, 1)
            result = _calculate_school_year()
            assert result == "2025-2026"


class TestGetSponsorshipTermsFromTier:
    """Tests for _get_sponsorship_terms_from_tier function."""

    def test_platinum_tier_terms(self):
        """Test platinum sponsor benefits are returned."""
        from src.tools.finance import _get_sponsorship_terms_from_tier

        result = _get_sponsorship_terms_from_tier("Platinum Sponsor")
        assert "Premier logo placement" in result
        assert "5-minute speaking slot" in result
        assert "resume book" in result

    def test_gold_tier_terms(self):
        """Test gold sponsor benefits are returned."""
        from src.tools.finance import _get_sponsorship_terms_from_tier

        result = _get_sponsorship_terms_from_tier("Gold Sponsor")
        assert "Booth space" in result
        assert "Social media feature posts (2x)" in result

    def test_unknown_tier_fallback(self):
        """Test unknown tier returns default message."""
        from src.tools.finance import _get_sponsorship_terms_from_tier

        result = _get_sponsorship_terms_from_tier("Unknown Tier")
        assert "Standard sponsorship benefits" in result

    def test_booth_sponsor_terms(self):
        """Test booth sponsor benefits are returned."""
        from src.tools.finance import _get_sponsorship_terms_from_tier

        result = _get_sponsorship_terms_from_tier("Booth Sponsor")
        assert "Booth presence" in result
        assert "interact with attendees" in result


class TestPopulateTemplateTags:
    """Tests for _populate_template_tags function."""

    @pytest.mark.asyncio
    async def test_basic_tag_population(self):
        """Test that all expected tags are populated."""
        from src.tools.finance import _populate_template_tags

        sponsor_data = {
            "company": "Test Corp",
            "contact_name": "John Doe",
            "tier": "Gold Sponsor",
            "amount": 15000,
        }
        event_config = {
            "name": "Blueprint",
            "year": 2026,
        }

        result = await _populate_template_tags(sponsor_data, event_config)

        # Verify all required tags are present
        assert "{{sponsor_company_name}}" in result
        assert "{{sponsor_name}}" in result
        assert "{{event_name}}" in result
        assert "{{year}}" in result
        assert "{{school_year}}" in result
        assert "{{sponsorship_amount}}" in result
        assert "{{sponsorship_terms}}" in result
        assert "{{invoice_due_date}}" in result
        assert "{{event_attendance_role}}" in result

    @pytest.mark.asyncio
    async def test_sponsor_data_populated_correctly(self):
        """Test sponsor data is correctly mapped to tags."""
        from src.tools.finance import _populate_template_tags

        sponsor_data = {
            "company": "Acme Inc",
            "contact_name": "Jane Smith",
            "tier": "Silver Sponsor",
            "amount": 5000,
        }
        event_config = {"name": "Hackathon", "year": 2026}

        result = await _populate_template_tags(sponsor_data, event_config)

        assert result["{{sponsor_company_name}}"] == "Acme Inc"
        assert result["{{sponsor_name}}"] == "Jane Smith"
        assert result["{{event_name}}"] == "Hackathon"
        assert result["{{sponsorship_amount}}"] == "$5,000.00"

    @pytest.mark.asyncio
    async def test_amount_fallback_to_tier_default(self):
        """Test amount uses tier default when not specified."""
        from src.tools.finance import _populate_template_tags

        sponsor_data = {
            "company": "No Amount Corp",
            "contact_name": "Test User",
            "tier": "Platinum Sponsor",
            "amount": None,  # No amount specified
        }
        event_config = {"name": "Event", "year": 2026}

        result = await _populate_template_tags(sponsor_data, event_config)

        # Platinum default is $25,000
        assert result["{{sponsorship_amount}}"] == "$25,000.00"

    @pytest.mark.asyncio
    async def test_invoice_due_date_30_days_out(self):
        """Test invoice due date is 30 days from now."""
        from src.tools.finance import _populate_template_tags

        sponsor_data = {"company": "Test", "contact_name": "Test", "tier": "TBD"}
        event_config = {"name": "Event", "year": 2026}

        with patch('src.tools.finance.datetime') as mock_datetime:
            mock_now = datetime(2026, 1, 15, 12, 0, 0)
            mock_datetime.now.return_value = mock_now
            mock_datetime.strftime = datetime.strftime
            
            # Need to also handle timedelta
            with patch('src.tools.finance.timedelta', side_effect=timedelta):
                result = await _populate_template_tags(sponsor_data, event_config)
                
                # 30 days from Jan 15 = Feb 14
                assert "February 14, 2026" in result["{{invoice_due_date}}"]

    @pytest.mark.asyncio
    async def test_attendance_role_inference_booth(self):
        """Test attendance role inferred from Booth tier."""
        from src.tools.finance import _populate_template_tags

        sponsor_data = {
            "company": "Booth Co",
            "contact_name": "Test",
            "tier": "Booth Sponsor",
        }
        event_config = {"name": "Event", "year": 2026}

        result = await _populate_template_tags(sponsor_data, event_config, None)

        assert result["{{event_attendance_role}}"] == "booth"

    @pytest.mark.asyncio
    async def test_attendance_role_override(self):
        """Test attendance role can be explicitly set."""
        from src.tools.finance import _populate_template_tags

        sponsor_data = {
            "company": "Override Co",
            "contact_name": "Test",
            "tier": "Gold Sponsor",
        }
        event_config = {"name": "Event", "year": 2026}

        result = await _populate_template_tags(
            sponsor_data, event_config, "mentor"
        )

        assert result["{{event_attendance_role}}"] == "mentor"


class TestReplaceTagsInDocument:
    """Tests for _replace_tags_in_document function."""

    def test_replace_simple_tag_in_paragraph(self):
        """Test replacing a simple tag in a paragraph."""
        from src.tools.finance import _replace_tags_in_document
        from docx import Document

        # Create a simple test document
        doc = Document()
        doc.add_paragraph("Hello {{sponsor_company_name}}!")

        tag_mapping = {
            "{{sponsor_company_name}}": "Test Corp"
        }

        _replace_tags_in_document(doc, tag_mapping)

        assert "Test Corp" in doc.paragraphs[0].text
        assert "{{sponsor_company_name}}" not in doc.paragraphs[0].text

    def test_replace_multiple_tags(self):
        """Test replacing multiple tags in same document."""
        from src.tools.finance import _replace_tags_in_document
        from docx import Document

        doc = Document()
        doc.add_paragraph("{{sponsor_name}} from {{sponsor_company_name}}")

        tag_mapping = {
            "{{sponsor_name}}": "John Doe",
            "{{sponsor_company_name}}": "Acme Inc"
        }

        _replace_tags_in_document(doc, tag_mapping)

        assert "John Doe" in doc.paragraphs[0].text
        assert "Acme Inc" in doc.paragraphs[0].text


# =============================================================================
# Integration Tests
# =============================================================================

class TestGenerateMouInvoice:
    """Integration tests for generate_mou_invoice tool."""

    @pytest.mark.asyncio
    async def test_generates_document_file(self):
        """Test that a document file is actually generated."""
        from src.tools.finance import generate_mou_invoice, OUTPUT_DIR, TEMPLATE_PATH

        # Skip if template doesn't exist
        if not TEMPLATE_PATH.exists():
            pytest.skip("Template file not found")

        # Mock database and org service
        with patch('src.tools.finance.org_service') as mock_org_service, \
             patch('src.tools.finance.db_service') as mock_db_service:
            
            # Setup mocks
            mock_config = MagicMock()
            mock_config.event.name = "Blueprint"
            mock_config.event.year = 2026
            mock_org_service.get_config = AsyncMock(return_value=mock_config)
            mock_db_service.find_one = AsyncMock(return_value=None)

            result = await generate_mou_invoice.ainvoke({
                "sponsor_company_name": "Test Corporation"
            })

            # Verify result contains expected content
            assert "[PENDING APPROVAL]" in result
            assert "MOU_Invoice_Test_Corporation" in result
            assert "Populated Values" in result

            # Verify file was created
            generated_files = list(OUTPUT_DIR.glob("MOU_Invoice_Test_Corporation*.docx"))
            assert len(generated_files) >= 1

            # Cleanup generated test file
            for f in generated_files:
                f.unlink()

    @pytest.mark.asyncio
    async def test_uses_mongodb_sponsor_data(self):
        """Test that MongoDB sponsor data is used when available."""
        from src.tools.finance import generate_mou_invoice, OUTPUT_DIR, TEMPLATE_PATH

        if not TEMPLATE_PATH.exists():
            pytest.skip("Template file not found")

        with patch('src.tools.finance.org_service') as mock_org_service, \
             patch('src.tools.finance.db_service') as mock_db_service:
            
            mock_config = MagicMock()
            mock_config.event.name = "Blueprint"
            mock_config.event.year = 2026
            mock_org_service.get_config = AsyncMock(return_value=mock_config)
            
            # Return sponsor data from "MongoDB"
            mock_db_service.find_one = AsyncMock(return_value={
                "company": "MongoDB Corp",
                "contact_name": "Database Dan",
                "tier": "Platinum Sponsor",
                "amount": 30000,
            })

            result = await generate_mou_invoice.ainvoke({
                "sponsor_company_name": "MongoDB"
            })

            # Verify MongoDB data was used
            assert "sponsor_company_name: MongoDB Corp" in result
            assert "sponsor_name: Database Dan" in result
            assert "$30,000.00" in result

            # Cleanup
            for f in OUTPUT_DIR.glob("MOU_Invoice_MongoDB*.docx"):
                f.unlink()

    @pytest.mark.asyncio
    async def test_template_missing_error(self):
        """Test error handling when template file is missing."""
        from src.tools.finance import generate_mou_invoice

        with patch('src.tools.finance.org_service') as mock_org_service, \
             patch('src.tools.finance.db_service') as mock_db_service, \
             patch('src.tools.finance.TEMPLATE_PATH') as mock_path:
            
            mock_config = MagicMock()
            mock_config.event.name = "Test"
            mock_config.event.year = 2026
            mock_org_service.get_config = AsyncMock(return_value=mock_config)
            mock_db_service.find_one = AsyncMock(return_value=None)
            mock_path.exists.return_value = False

            result = await generate_mou_invoice.ainvoke({
                "sponsor_company_name": "Missing Template Corp"
            })

            assert "Error" in result or "not found" in result


# =============================================================================
# Template Tag Verification Tests
# =============================================================================

class TestTemplateTags:
    """Tests to verify all template tags are handled."""

    def test_all_expected_tags_defined(self):
        """Verify all expected tags have mappings defined."""
        expected_tags = [
            "{{sponsor_company_name}}",
            "{{sponsor_name}}",
            "{{event_name}}",
            "{{year}}",
            "{{school_year}}",
            "{{sponsorship_amount}}",
            "{{sponsorship_terms}}",
            "{{invoice_due_date}}",
            "{{event_attendance_role}}",
        ]

        @pytest.mark.asyncio
        async def verify_tags():
            from src.tools.finance import _populate_template_tags

            sponsor_data = {
                "company": "Test",
                "contact_name": "Test",
                "tier": "Gold Sponsor",
            }
            event_config = {"name": "Event", "year": 2026}

            result = await _populate_template_tags(sponsor_data, event_config)

            for tag in expected_tags:
                assert tag in result, f"Missing tag mapping: {tag}"
                assert result[tag], f"Empty value for tag: {tag}"

        import asyncio
        asyncio.run(verify_tags())

    def test_no_unpopulated_tags_in_output(self):
        """Verify generated documents have no remaining {{...}} tags."""
        from src.tools.finance import _replace_tags_in_document, _populate_template_tags
        from docx import Document
        import asyncio

        # Create document with all tags
        doc = Document()
        doc.add_paragraph("{{sponsor_company_name}} - {{sponsor_name}}")
        doc.add_paragraph("Event: {{event_name}} {{year}} ({{school_year}})")
        doc.add_paragraph("Amount: {{sponsorship_amount}}")
        doc.add_paragraph("Terms: {{sponsorship_terms}}")
        doc.add_paragraph("Due: {{invoice_due_date}}")
        doc.add_paragraph("Role: {{event_attendance_role}}")

        async def verify():
            sponsor_data = {
                "company": "Complete Corp",
                "contact_name": "Full Coverage",
                "tier": "Gold Sponsor",
                "amount": 15000,
            }
            event_config = {"name": "Blueprint", "year": 2026}
            
            tag_mapping = await _populate_template_tags(sponsor_data, event_config)
            _replace_tags_in_document(doc, tag_mapping)

            # Check no remaining template tags
            full_text = "\n".join([p.text for p in doc.paragraphs])
            remaining_tags = re.findall(r'\{\{[^}]+\}\}', full_text)
            assert not remaining_tags, f"Unpopulated tags found: {remaining_tags}"

        asyncio.run(verify())
