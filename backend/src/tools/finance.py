"""
Finance tools for Nexus.

Provides tools for budget management, invoicing, and financial tracking.
"""

from langchain_core.tools import tool

# Mock budget data
MOCK_BUDGET = {
    "total_budget": 100000,
    "confirmed_sponsorship": 65000,
    "pending_sponsorship": 15000,
    "expenses": {
        "venue": 15000,
        "catering": 8000,
        "marketing": 5000,
        "swag": 3000,
        "prizes": 10000,
        "misc": 2000,
    },
    "remaining": 22000,
}


@tool
async def check_budget_status(category: str = "all") -> str:
    """Check the current budget status.

    Args:
        category: Specific category or 'all' for full overview

    Returns:
        Budget status summary
    """
    if category == "all":
        expenses_list = "\n".join([
            f"  - {k.title()}: ${v:,}"
            for k, v in MOCK_BUDGET["expenses"].items()
        ])
        total_expenses = sum(MOCK_BUDGET["expenses"].values())

        return f"""Budget Overview:

Income:
- Confirmed Sponsorship: ${MOCK_BUDGET['confirmed_sponsorship']:,}
- Pending Sponsorship: ${MOCK_BUDGET['pending_sponsorship']:,}
- Total Target: ${MOCK_BUDGET['total_budget']:,}

Expenses:
{expenses_list}
- Total Expenses: ${total_expenses:,}

Balance:
- Available: ${MOCK_BUDGET['confirmed_sponsorship'] - total_expenses:,}
- Pending + Available: ${MOCK_BUDGET['confirmed_sponsorship'] + MOCK_BUDGET['pending_sponsorship'] - total_expenses:,}"""
    else:
        amount = MOCK_BUDGET["expenses"].get(category.lower(), 0)
        return f"{category.title()} Budget: ${amount:,}"


@tool
async def update_budget_sheet(
    category: str,
    amount: float,
    description: str,
    transaction_type: str = "expense"
) -> str:
    """Log a budget transaction.

    Args:
        category: Budget category
        amount: Transaction amount
        description: Description of the transaction
        transaction_type: 'expense' or 'income'

    Returns:
        Confirmation of logged transaction
    """
    return f"""Budget Transaction Logged:

Type: {transaction_type.title()}
Category: {category.title()}
Amount: ${amount:,.2f}
Description: {description}

Budget sheet has been updated."""


@tool
async def draft_mou(
    sponsor_name: str,
    tier: str,
    amount: float,
    benefits: str
) -> str:
    """Draft a Memorandum of Understanding for a sponsor (requires approval).

    Args:
        sponsor_name: Name of the sponsor company
        tier: Sponsorship tier (Platinum, Gold, Silver, etc.)
        amount: Sponsorship amount
        benefits: Key benefits included

    Returns:
        Draft MOU content
    """
    return f"""[PENDING APPROVAL] Draft MOU:

MEMORANDUM OF UNDERSTANDING
Between: Blueprint Event Team and {sponsor_name}

Sponsorship Details:
- Tier: {tier}
- Amount: ${amount:,.2f}
- Event: Blueprint 2024

Benefits Included:
{benefits}

Terms:
1. Payment due within 30 days of signing
2. Logo placement on all event materials
3. Recognition in opening/closing ceremonies
4. Access to attendee resume book (with consent)

This MOU requires your approval before sending to the sponsor."""


@tool
async def generate_invoice(
    sponsor_name: str,
    contact_email: str,
    amount: float,
    tier: str,
    due_date: str
) -> str:
    """Generate an invoice for a sponsor (requires approval).

    Args:
        sponsor_name: Sponsor company name
        contact_email: Billing contact email
        amount: Invoice amount
        tier: Sponsorship tier
        due_date: Payment due date

    Returns:
        Invoice details
    """
    invoice_number = f"INV-2024-{hash(sponsor_name) % 10000:04d}"

    return f"""[PENDING APPROVAL] Invoice Generated:

Invoice #: {invoice_number}
Bill To: {sponsor_name}
Contact: {contact_email}

Description: {tier} Sponsorship - Blueprint 2024
Amount: ${amount:,.2f}
Due Date: {due_date}

Payment Methods:
- Bank Transfer (details to be provided)
- Check (payable to Blueprint Events)

This invoice requires your approval before sending."""


@tool
async def get_sponsorship_financials(tier: str = "all") -> str:
    """Get financial summary for sponsorships.

    Args:
        tier: Specific tier or 'all'

    Returns:
        Financial summary
    """
    tiers = {
        "Platinum": {"count": 1, "amount": 25000, "total": 25000},
        "Gold": {"count": 2, "amount": 15000, "total": 30000},
        "Silver": {"count": 2, "amount": 5000, "total": 10000},
    }

    if tier != "all" and tier in tiers:
        t = tiers[tier]
        return f"""{tier} Sponsorship:
- Partners: {t['count']}
- Per Partner: ${t['amount']:,}
- Total: ${t['total']:,}"""

    summary = "\n".join([
        f"- {name}: {data['count']} sponsors = ${data['total']:,}"
        for name, data in tiers.items()
    ])
    total = sum(t["total"] for t in tiers.values())

    return f"""Sponsorship Financial Summary:

{summary}

Total Confirmed: ${total:,}
Goal: $100,000
Progress: {int(total/100000*100)}%"""


# =============================================================================
# MOU and Invoice Document Generation
# =============================================================================

import logging
import os
import re
from datetime import datetime, timedelta
from pathlib import Path
from typing import Optional

from docx import Document

from src.services.database import db_service
from src.services.organization import org_service, DEFAULT_ORG_ID

logger = logging.getLogger(__name__)

# Path to the template file
TEMPLATE_PATH = Path(__file__).parent / "templates" / "Draft MOU and invoice document.docx"
OUTPUT_DIR = Path(__file__).parent / "generated"


def _calculate_school_year() -> str:
    """
    Get academic year string based on current date.
    
    Returns:
        str: School year in format "2025-2026"
    """
    now = datetime.now()
    # Academic year typically starts in September
    # If we're in Jan-Aug, we're in the second half of the school year
    if now.month < 9:
        start_year = now.year - 1
    else:
        start_year = now.year
    return f"{start_year}-{start_year + 1}"


def _get_sponsorship_terms_from_tier(tier: str) -> str:
    """
    Get sponsorship terms/benefits based on tier.
    
    Args:
        tier: Sponsorship tier name
        
    Returns:
        str: Formatted sponsorship terms
    """
    tier_benefits = {
        "Platinum Sponsor": """- Premier logo placement on all event materials
- Dedicated booth space in prime location
- 5-minute speaking slot during opening ceremony
- Access to attendee resume book (with consent)
- Social media feature posts (3x)
- Company swag distribution to all attendees""",
        "Platinum": """- Premier logo placement on all event materials
- Dedicated booth space in prime location
- 5-minute speaking slot during opening ceremony
- Access to attendee resume book (with consent)
- Social media feature posts (3x)
- Company swag distribution to all attendees""",
        "Gold Sponsor": """- Logo placement on all event materials
- Booth space in venue
- Recognition during opening ceremony
- Access to attendee resume book (with consent)
- Social media feature posts (2x)""",
        "Gold": """- Logo placement on all event materials
- Booth space in venue
- Recognition during opening ceremony
- Access to attendee resume book (with consent)
- Social media feature posts (2x)""",
        "Silver Sponsor": """- Logo on event website and materials
- Recognition during ceremonies
- Access to attendee resume book (with consent)
- Social media mention""",
        "Silver": """- Logo on event website and materials
- Recognition during ceremonies
- Access to attendee resume book (with consent)
- Social media mention""",
        "Bronze Sponsor": """- Logo on event website
- Recognition during ceremonies
- Social media mention""",
        "Bronze": """- Logo on event website
- Recognition during ceremonies
- Social media mention""",
        "Booth Sponsor": """- Booth presence at event
- Logo on event materials
- Opportunity to interact with attendees""",
        "Booth": """- Booth presence at event
- Logo on event materials
- Opportunity to interact with attendees""",
        "In-Kind Sponsor": """- Logo on event materials
- Recognition for contribution
- Social media mention""",
        "In-Kind": """- Logo on event materials
- Recognition for contribution
- Social media mention""",
    }
    return tier_benefits.get(tier, "- Standard sponsorship benefits as agreed")


async def _populate_template_tags(
    sponsor_data: dict,
    event_config: dict,
    attendance_role: Optional[str] = None,
) -> dict:
    """
    Build tag replacement mapping from sponsor and event data.
    
    Args:
        sponsor_data: Dict containing sponsor information from MongoDB
        event_config: Dict containing event configuration
        attendance_role: Optional role for event attendance
        
    Returns:
        dict: Mapping of template tags to their values
    """
    now = datetime.now()
    
    # Calculate invoice due date (30 days from now)
    due_date = now + timedelta(days=30)
    
    # Get tier first (needed for both amount lookup and terms)
    tier = sponsor_data.get("tier", "TBD")

    # Tier amount defaults
    tier_amounts = {
        "Platinum Sponsor": 25000,
        "Platinum": 25000,
        "Gold Sponsor": 15000,
        "Gold": 15000,
        "Silver Sponsor": 5000,
        "Silver": 5000,
        "Bronze Sponsor": 2500,
        "Bronze": 2500,
        "Booth Sponsor": 1500,
        "Booth": 1500,
        "In-Kind Sponsor": 0,
        "In-Kind": 0,
    }

    # Get sponsorship amount - use provided amount, otherwise use tier default
    amount = sponsor_data.get("amount")
    if not amount:
        amount = tier_amounts.get(tier, 0)
    terms = _get_sponsorship_terms_from_tier(tier)
    
    # Determine attendance role
    role = attendance_role
    if not role:
        # Infer from tier if not provided
        if "Booth" in tier:
            role = "booth"
        elif tier == "TBD":
            role = "networking delegate"
        else:
            role = "sponsor representative"
    
    return {
        "{{sponsor_company_name}}": sponsor_data.get("company", ""),
        "{{sponsor_name}}": sponsor_data.get("contact_name", ""),
        "{{event_name}}": event_config.get("name", "Event"),
        "{{year}}": str(now.year),
        "{{school_year}}": _calculate_school_year(),
        "{{sponsorship_amount}}": f"${amount:,.2f}" if amount else "TBD",
        "{{sponsorship_terms}}": terms,
        "{{invoice_due_date}}": due_date.strftime("%B %d, %Y"),
        "{{event_attendance_role}}": role,
    }


def _replace_tags_in_paragraph(paragraph, tag_mapping: dict) -> None:
    """
    Replace template tags in a paragraph, handling tags split across runs.

    Word documents split text into "runs" based on formatting. A tag like
    {{sponsor_company_name}} might be split across multiple runs. This function
    combines all run text, performs replacements, and rebuilds with first run's formatting.

    Args:
        paragraph: python-docx Paragraph object
        tag_mapping: Dict of tags to their replacement values
    """
    # Get full paragraph text
    full_text = paragraph.text

    # Check if any tags exist in this paragraph
    has_tags = any(tag in full_text for tag in tag_mapping.keys())
    if not has_tags:
        return

    # Perform all replacements on the full text
    for tag, value in tag_mapping.items():
        full_text = full_text.replace(tag, value)

    # Clear existing runs and set new text
    # Preserve formatting from first run if it exists
    if paragraph.runs:
        # Store first run's formatting
        first_run = paragraph.runs[0]

        # Clear all runs by setting text to empty
        for run in paragraph.runs:
            run.text = ""

        # Set the replaced text in the first run
        first_run.text = full_text
    else:
        # No runs exist, just set the text directly
        paragraph.text = full_text


def _replace_tags_in_document(doc: Document, tag_mapping: dict) -> None:
    """
    Replace all template tags in a Word document.

    Args:
        doc: python-docx Document object
        tag_mapping: Dict of tags to their replacement values
    """
    # Replace in paragraphs
    for paragraph in doc.paragraphs:
        _replace_tags_in_paragraph(paragraph, tag_mapping)

    # Replace in tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _replace_tags_in_paragraph(paragraph, tag_mapping)


@tool
async def generate_mou_invoice(
    sponsor_company_name: str,
    tier: Optional[str] = None,
    contact_name: Optional[str] = None,
    contact_email: Optional[str] = None,
    amount: Optional[float] = None,
    event_name: Optional[str] = None,
    attendance_role: Optional[str] = None,
) -> str:
    """
    Generate MOU and invoice document from template (requires approval).
    
    Populates the template with provided sponsor data and calculated values
    like current year, school year, and invoice due date.
    
    Args:
        sponsor_company_name: Company name of the sponsor
        tier: Sponsorship tier (Platinum, Gold, Silver, Bronze, Booth, In-Kind)
        contact_name: Name of the contact person at the sponsor company
        contact_email: Email address of the contact person
        amount: Sponsorship amount in dollars (if different from tier default)
        event_name: Event name (uses org default if not provided)
        attendance_role: Optional role (e.g., 'booth', 'mentor', 'networking delegate')
    
    Returns:
        str: Path to generated document and preview of populated values
    """
    try:
        # Ensure output directory exists
        OUTPUT_DIR.mkdir(exist_ok=True)
        
        # Get organization config for event details
        org_config = await org_service.get_config(DEFAULT_ORG_ID)
        
        # Get event configuration
        event_cfg = {
            "name": event_name or org_config.event.name,
            "year": org_config.event.year,
        }
        
        # Try to look up sponsor in MongoDB for additional info
        sponsor_data = None
        try:
            collection = db_service.db["sponsor_partnerships"]
            result = await collection.find_one(
                {"company": {"$regex": sponsor_company_name, "$options": "i"}}
            )
            if result:
                sponsor_data = result
        except Exception as e:
            logger.warning(f"Could not fetch sponsor from MongoDB: {e}")
        
        # Use provided values, fall back to MongoDB data, then defaults
        final_sponsor_data = {
            "company": sponsor_company_name,
            "contact_name": contact_name or (sponsor_data.get("contact_name") if sponsor_data else "") or "",
            "contact_email": contact_email or (sponsor_data.get("contact_email") if sponsor_data else "") or "",
            "tier": tier or (sponsor_data.get("tier") if sponsor_data else None) or "TBD",
            "amount": amount or (sponsor_data.get("amount") if sponsor_data else None),
        }
        
        # Build tag mapping
        tag_mapping = await _populate_template_tags(
            final_sponsor_data, 
            event_cfg, 
            attendance_role
        )
        
        # Check if template exists
        if not TEMPLATE_PATH.exists():
            return f"Error: Template file not found at {TEMPLATE_PATH}"
        
        # Load and process template
        doc = Document(str(TEMPLATE_PATH))
        _replace_tags_in_document(doc, tag_mapping)
        
        # Generate output filename
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        safe_name = re.sub(r'[^\w\s-]', '', sponsor_company_name).replace(' ', '_')
        output_filename = f"MOU_Invoice_{safe_name}_{timestamp}.docx"
        output_path = OUTPUT_DIR / output_filename
        
        # Save document
        doc.save(str(output_path))
        
        # Build download URL
        download_url = f"/api/files/{output_filename}"

        # Build preview of populated values
        preview_lines = ["[PENDING APPROVAL] MOU/Invoice Document Generated:", ""]
        preview_lines.append(f"📄 **File:** {output_filename}")
        preview_lines.append(f"📥 **Download:** [{output_filename}]({download_url})")
        preview_lines.append("")
        preview_lines.append("**Populated Values:**")
        for tag, value in tag_mapping.items():
            tag_name = tag.replace("{{", "").replace("}}", "")
            # Truncate long values for preview
            display_value = value if len(value) < 100 else value[:100] + "..."
            preview_lines.append(f"  • {tag_name}: {display_value}")

        preview_lines.append("")
        preview_lines.append("⚠️ This document requires your approval before sending to the sponsor.")

        # Return structured response with download URL for frontend
        return f"DOWNLOAD_URL:{download_url}\n" + "\n".join(preview_lines)
        
    except Exception as e:
        logger.error(f"Error generating MOU/Invoice: {e}")
        return f"Error generating document: {str(e)}"


FINANCE_TOOLS = [
    check_budget_status,
    update_budget_sheet,
    draft_mou,
    generate_invoice,
    get_sponsorship_financials,
    generate_mou_invoice,
]

